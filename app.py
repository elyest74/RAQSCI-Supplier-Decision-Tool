import streamlit as st
import pandas as pd
import plotly.graph_objects as go
import io, tempfile

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4

# ================= CONFIG =================
st.set_page_config(layout="wide")

SECTORES = [
    "Industrial / Manufacturing","Energía / Utilities","Construcción / Ingeniería",
    "Automoción","Aeroespacial / Defensa","Alimentación / FMCG",
    "Retail / Moda / eCommerce","Farmacéutico / Healthcare",
    "Tecnología / Electrónica","Logística / Transporte",
    "Servicios / Outsourcing","Químico / Materias primas","Otros"
]

mapping = {"Bajo (1)":1,"Medio (3)":3,"Alto (5)":5}

# ================= PESOS =================
def get_weights(sector):
    s = sector.lower()
    if "retail" in s or "moda" in s:
        return {"R":0.05,"A":0.20,"Q":0.30,"S":0.20,"C":0.15,"I":0.10}
    if "industrial" in s:
        return {"R":0.10,"A":0.35,"Q":0.25,"S":0.10,"C":0.15,"I":0.05}
    if "farmacéutico" in s:
        return {"R":0.30,"A":0.25,"Q":0.25,"S":0.10,"C":0.05,"I":0.05}
    return {"R":0.10,"A":0.25,"Q":0.20,"S":0.15,"C":0.20,"I":0.10}

# ================= EXCEL =================
def generate_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "RAQSCI_INPUT"

    headers = [
        "Proveedor","Categoria","Subcategoria",
        "R_certificaciones","R_cumplimiento","R_ESG",
        "A_capacidad","A_dependencia","A_resiliencia",
        "Q_defectos","Q_consistencia","Q_auditorias",
        "S_leadtime","S_flexibilidad","S_soporte",
        "C_precio","C_logistica","C_inventario","C_TCO",
        "I_mejora","I_ID","I_digitalizacion"
    ]

    ws.append(headers)

    for i in range(1,len(headers)+1):
        ws.column_dimensions[get_column_letter(i)].width = 22

    dv = DataValidation(type="list", formula1='"Bajo (1),Medio (3),Alto (5)"')
    ws.add_data_validation(dv)

    for r in range(2,200):
        for c in range(4,len(headers)+1):
            dv.add(ws.cell(row=r,column=c))

    ws.freeze_panes="D2"

    ws2 = wb.create_sheet("COMO_PUNTUAR")

    guide = [
        ["CRITERIO","CÓMO MEDIR","BAJO","MEDIO","ALTO"],
        ["Defectos","% defectos",">5%","~2%","<0.5%"],
        ["Lead time","semanas",">8","~4","<2"],
        ["Dependencia","nº proveedores","1","2","3+"],
        ["Precio","vs mercado","alto","medio","bajo"],
        ["Flexibilidad","cambios","baja","media","alta"],
        ["Innovación","propuestas","nula","media","alta"]
    ]

    for r in guide:
        ws2.append(r)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

# ================= MODELO =================
def calc_raqsci(r,a,q,s,c,i,sector):
    w = get_weights(sector)
    score = r*w["R"]+a*w["A"]+q*w["Q"]+s*w["S"]+c*w["C"]+i*w["I"]
    status = "APTO" if r>=3 and q>=3 and a>=3 else "NO APTO"
    return round(score,2), status

def calc_kraljic(q,c,a):
    impacto=(q+c)/2
    riesgo=a
    if impacto>=4 and riesgo>=4: return impacto,riesgo,"Estratégica"
    elif impacto>=4: return impacto,riesgo,"Apalancada"
    elif riesgo>=4: return impacto,riesgo,"Cuello botella"
    return impacto,riesgo,"No crítico"

# ================= PROCESS =================
def process(df,sector):

    for col in df.columns[3:]:
        if df[col].isnull().any():
            st.error(f"Error: valores vacíos en {col}")
            st.stop()

    for col in df.columns[3:]:
        df[col] = df[col].map(mapping)

    data=[]
    for _,row in df.iterrows():

        r=row[3:6].mean()
        a=row[6:9].mean()
        q=row[9:12].mean()
        s=row[12:15].mean()
        c=row[15:19].mean()
        i=row[19:22].mean()

        score,status = calc_raqsci(r,a,q,s,c,i,sector)
        impacto,riesgo,k = calc_kraljic(q,c,a)

        data.append({
            "Proveedor":row["Proveedor"],
            "Score":score,
            "Estado":status,
            "Impacto":impacto,
            "Riesgo":riesgo,
            "Kraljic":k,
            "C":c,"Q":q
        })

    return pd.DataFrame(data)

# ================= MATRIZ =================
def plot_matrix(df):
    fig=go.Figure()

    colors={
        "Estratégica":"#E74C3C",
        "Apalancada":"#2ECC71",
        "Cuello botella":"#F1C40F",
        "No crítico":"#3498DB"
    }

    for cat in df["Kraljic"].unique():
        sub=df[df["Kraljic"]==cat]
        fig.add_trace(go.Scatter(
            x=sub["Impacto"], y=sub["Riesgo"],
            mode="markers",
            name=cat,
            marker=dict(size=12,color=colors[cat]),
            text=sub["Proveedor"],
            hovertemplate="<b>%{text}</b>"
        ))

    fig.add_shape(type="line",x0=3,x1=3,y0=0,y1=5)
    fig.add_shape(type="line",x0=0,x1=5,y0=3,y1=3)

    fig.update_layout(legend=dict(x=1.05,y=1))
    return fig

# ================= PDF =================
def generate_pdf(df,fig):
    styles=getSampleStyleSheet()
    buffer=tempfile.NamedTemporaryFile(delete=False,suffix=".pdf")
    doc=SimpleDocTemplate(buffer.name,pagesize=A4)

    elems=[]
    top=df.sort_values(by="Score",ascending=False).iloc[0]

    elems.append(Paragraph("Strategic Supplier Decision Tool",styles["Title"]))
    elems.append(PageBreak())

    elems.append(Paragraph("Executive Summary",styles["Heading2"]))
    elems.append(Paragraph(f"Proveedor recomendado: {top['Proveedor']}",styles["Normal"]))
    elems.append(PageBreak())

    img=tempfile.NamedTemporaryFile(delete=False,suffix=".png").name
    fig.write_image(img)
    elems.append(Image(img,width=450,height=300))

    doc.build(elems)
    return open(buffer.name,"rb").read()

# ================= UI =================
st.markdown("""
<div style='background:#1F3A5F;padding:20px'>
<h1 style='color:white'>Strategic Supplier Decision Tool</h1>
<p style='color:white'>Integrated Kraljic & RAQSCI Analysis for Procurement Leaders</p>
</div>
""", unsafe_allow_html=True)

sector = st.selectbox("Sector",SECTORES)

with st.sidebar:
    st.image("elymar.png", width=120)
    st.download_button("📥 Descargar plantilla", generate_excel(), "plantilla.xlsx")
    file = st.file_uploader("📤 Upload Excel", type=["xlsx"])

if file:
    df = process(pd.read_excel(file),sector)

    # KPIs
    c1,c2,c3=st.columns(3)
    c1.metric("Proveedores",len(df))
    c2.metric("% NO APTO",round((df["Estado"]=="NO APTO").mean()*100,1))
    c3.metric("Score medio",round(df["Score"].mean(),2))

    # matriz + ranking
    col1,col2=st.columns([2,1])

    with col1:
        fig=plot_matrix(df)
        st.plotly_chart(fig,use_container_width=True)

    with col2:
        st.dataframe(df.sort_values(by="Score",ascending=False))

    # insights
    st.subheader("Insights")
    if df["Riesgo"].mean()>4:
        st.warning("Riesgo alto")
    if (df["C"]>4).any() and (df["Q"]<2).any():
        st.warning("Incoherencia coste-calidad")

    # recomendación
    top=df[df["Estado"]=="APTO"].sort_values(by="Score",ascending=False).iloc[0]
    st.success(f"Recomendación: {top['Proveedor']}")

    # PDF
    pdf=generate_pdf(df,fig)
    st.download_button("📄 Descargar PDF",pdf,"report.pdf")
